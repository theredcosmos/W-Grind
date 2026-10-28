"""
Import the DSA master workbook into the app's static data files.

Usage:
    python tools/import_xlsx.py <path-to-xlsx>

Outputs (all regenerated, idempotent):
    app/src/data/questions.json   one record per unique LeetCode problem (keyed by URL slug)
    app/src/data/lists.json       ordered Blind 75 / NeetCode 150 / Striver SDE sheets
    app/src/data/manual.json      ADHD manual, pattern decoder and element-linker guide
    app/src/data/legacyIds.json   old numeric question id -> new slug (progress migration)

Curated lists that are not in the workbook live in tools/data/*.json (e.g. LeetCode's
Top Interview 150); problems they add get numbered after the workbook's last Q#.

Natural key = LeetCode slug parsed from the URL, so a problem that appears in two
sections of the master sheet is stored once and its list memberships are merged.
"""
import json
import os
import re
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(ROOT, "app", "src", "data")

MASTER = "1_Master_562_Tracker"
LIST_SHEETS = {
    "blind75": "3_Blind_75_Sheet",
    "neetcode150": "4_NeetCode_150_Sheet",
    "striver": "5_Striver_SDE_180_Sheet",
}
LIST_META = {
    "blind75": {"name": "Blind 75", "short": "B75"},
    "neetcode150": {"name": "NeetCode 150", "short": "NC150"},
    "striver": {"name": "Striver SDE Sheet", "short": "SDE"},
}
DIFFICULTIES = {"Easy", "Medium", "Hard"}
CURATED_FILES = {"top150": os.path.join(ROOT, "tools", "data", "top_interview_150.json")}


def slug_of(url):
    """Canonical natural key: the /problems/<slug>/ segment, lowercased."""
    m = re.search(r"leetcode\.com/problems/([^/?#]+)", url or "")
    if not m:
        raise ValueError(f"Not a LeetCode problem URL: {url!r}")
    return m.group(1).strip().lower()


def clean(v):
    if v is None:
        return ""
    return str(v).strip()


def strip_prefix(label):
    """'Section 1: Arrays' -> 'Arrays', '1.1 Two Pointers' -> 'Two Pointers'."""
    label = re.sub(r"^Section\s+\d+:\s*", "", label)
    return re.sub(r"^\d+(\.\d+)*\s+", "", label)


def load_master(wb):
    ws = wb[MASTER]
    questions = {}
    by_master_no = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        q_no, _, _, in_blind, in_nc, in_striver, week, section, sub, pattern, title, url, diff = row[:13]
        trigger, linker = row[19], row[20]
        slug = slug_of(url)
        by_master_no[int(q_no)] = slug
        if diff not in DIFFICULTIES:
            raise ValueError(f"Q{q_no}: bad difficulty {diff!r}")
        if not clean(title):
            raise ValueError(f"Q{q_no}: missing title")

        if slug in questions:
            # Duplicate row for the same problem: keep first placement, remember the alias.
            questions[slug]["also"].append(strip_prefix(clean(sub)))
            continue

        questions[slug] = {
            "id": slug,
            "n": int(q_no),
            "title": clean(title),
            "url": f"https://leetcode.com/problems/{slug}/",
            "difficulty": diff,
            "week": clean(week),
            "section": strip_prefix(clean(section)),
            "topic": strip_prefix(clean(sub)),
            "pattern": clean(pattern),
            "trigger": clean(trigger),
            "linker": clean(linker),
            "lists": [],
            "also": [],
        }
        # Master flags are informational; list sheets below are the source of truth.
        _ = (in_blind, in_nc, in_striver)
    return questions, by_master_no


def load_lists(wb, questions, by_master_no):
    lists = {}
    for key, sheet in LIST_SHEETS.items():
        ws = wb[sheet]
        items, seen = [], set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            master_no, category, url = row[1], clean(row[4]), row[7]
            slug = slug_of(url) if url else by_master_no[int(master_no)]
            if slug not in questions:
                raise ValueError(f"{sheet}: {slug} not found in master")
            if slug in seen:
                continue
            seen.add(slug)
            items.append({"id": slug, "group": category})
            if key not in questions[slug]["lists"]:
                questions[slug]["lists"].append(key)
        lists[key] = {**LIST_META[key], "items": items}
    return lists


def load_curated(questions, lists):
    """Merge JSON-defined lists; problems missing from the workbook come from the file's 'extra' map."""
    next_n = max(q["n"] for q in questions.values()) + 1
    for key, path in CURATED_FILES.items():
        with open(path, encoding="utf-8") as f:
            spec = json.load(f)
        items = []
        for group in spec["groups"]:
            for slug in group["slugs"]:
                if slug not in questions:
                    extra = spec["extra"].get(slug)
                    if not extra:
                        raise ValueError(f"{path}: {slug} is not in the workbook and has no 'extra' entry")
                    if extra["difficulty"] not in DIFFICULTIES:
                        raise ValueError(f"{path}: {slug} has bad difficulty {extra['difficulty']!r}")
                    questions[slug] = {
                        "id": slug, "n": next_n, "url": f"https://leetcode.com/problems/{slug}/",
                        "week": "", **extra, "lists": [], "also": [],
                    }
                    next_n += 1
                items.append({"id": slug, "group": group["group"]})
                if key not in questions[slug]["lists"]:
                    questions[slug]["lists"].append(key)
        lists[key] = {"name": spec["name"], "short": spec["short"], "items": items}


def load_manual(wb):
    def table(sheet, header_row):
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        title = clean(rows[0][0])
        headers = [clean(h) for h in rows[header_row]]
        body = [[clean(c) for c in r] for r in rows[header_row + 1:] if any(r)]
        return {"title": title, "headers": headers, "rows": body}

    return {
        "antiBlank": table("6_ADHD_Anti_Blank_Manual", 1),
        "decoder": table("7_Pattern_Recognition_Decoder", 1),
        "linker": table("8_DSA_Element_Linker_Guide", 1),
    }


def legacy_map(questions):
    """Map ids from the previous questions.json (LeetCode numbers) onto new slugs."""
    path = os.path.join(DATA_DIR, "questions.json")
    if not os.path.exists(path):
        return {}
    with open(path, encoding="utf-8") as f:
        old = json.load(f)
    if old and isinstance(old[0].get("id"), str):
        # Already the new format; keep the existing map.
        existing = os.path.join(DATA_DIR, "legacyIds.json")
        if os.path.exists(existing):
            with open(existing, encoding="utf-8") as f:
                return json.load(f)
        return {}
    out = {}
    for q in old:
        try:
            slug = slug_of(q.get("url"))
        except ValueError:
            continue
        if slug in questions:
            out[str(q["id"])] = slug
    return out


def write(name, data):
    path = os.path.join(DATA_DIR, name)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
    print(f"wrote {path} ({os.path.getsize(path) // 1024} KB)")


def main():
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True, read_only=True)
    questions, by_master_no = load_master(wb)
    lists = load_lists(wb, questions, by_master_no)
    load_curated(questions, lists)
    legacy = legacy_map(questions)

    ordered = sorted(questions.values(), key=lambda q: q["n"])
    for q in ordered:
        if not q["also"]:
            del q["also"]

    write("legacyIds.json", legacy)
    write("questions.json", ordered)
    write("lists.json", lists)
    write("manual.json", load_manual(wb))

    print(f"{len(ordered)} unique problems; "
          + ", ".join(f"{v['name']}: {len(v['items'])}" for v in lists.values())
          + f"; {len(legacy)} legacy ids mapped")


if __name__ == "__main__":
    main()
